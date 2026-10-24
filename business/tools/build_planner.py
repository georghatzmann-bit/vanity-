#!/usr/bin/env python3
"""Baut den Finanzplaner als Excel-Datei.

Erzeugt zwei Dateien in business/produkt/:
  Finanzplaner-<Jahr>.xlsx       -> die Verkaufsversion (leer, mit je einer Beispielzeile)
  Finanzplaner-<Jahr>-DEMO.xlsx  -> mit Beispieldaten, fuer Screenshots und Werbung

Aufruf:  python3 build_planner.py [jahr]
Danach:  Formeln mit LibreOffice berechnen lassen (siehe README im Ordner business/).
"""
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

YEAR = int(sys.argv[1]) if len(sys.argv) > 1 else 2027
OUT = Path(__file__).resolve().parent.parent / "produkt"

FONT = "Arial"
ACCENT = "1F5F5B"
ACCENT_LIGHT = "E3EFEE"
INPUT = "FFF7D6"
INCOME = "2E7D32"
INCOME_LIGHT = "E8F5E9"
EXPENSE = "B42318"
EXPENSE_LIGHT = "FDECEA"
GREY = "6B7280"
LINE = "D1D5DB"

EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
EUR0 = '#,##0 "€";-#,##0 "€";"-"'
PCT = '0.0%;-0.0%;"-"'
DATE = "DD.MM.YYYY"

MONTHS = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]

# Kategorien-Tab: Zeilen 5-12 Einnahmen, 13-37 Ausgaben (eine zusammenhaengende Liste fuer das Dropdown)
INC_FIRST, INC_LAST = 5, 12
EXP_FIRST, EXP_LAST = 13, 37
INCOME_CATS = ["Gehalt / Lohn", "Nebenjob", "Verkäufe (Kleinanzeigen etc.)", "Sonstige Einnahmen"]
EXPENSE_CATS = [
    "Miete & Nebenkosten", "Strom & Internet", "Lebensmittel", "Handy & Abos",
    "Mobilität (Auto/ÖPNV)", "Versicherungen", "Freizeit & Ausgehen", "Kleidung",
    "Gesundheit & Drogerie", "Geschenke", "Urlaub", "Sonstiges",
]

# Buchungen-Tab
B_FIRST, B_LAST = 5, 1004
BK_DATE = f"Buchungen!$A${B_FIRST}:$A${B_LAST}"
BK_CAT = f"Buchungen!$B${B_FIRST}:$B${B_LAST}"
BK_AMT = f"Buchungen!$D${B_FIRST}:$D${B_LAST}"

# Monatsuebersicht-Tab
MO_INC_FIRST, MO_INC_LAST, MO_INC_SUM = 6, 13, 14
MO_EXP_FIRST, MO_EXP_LAST, MO_EXP_SUM = 17, 41, 42
MO_NET, MO_RATE, MO_BAL, MO_ACTIVE = 44, 45, 46, 48

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)


def font(size=10, bold=False, color="111827", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = font(18, True, ACCENT)
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = font(10, color=GREY, italic=True)


def header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def style_all(ws, max_row, max_col):
    """Setzt Arial auf alle Zellen, die noch die Standardschrift haben."""
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = font()


# --------------------------------------------------------------------------- Start
def build_start(wb, demo):
    ws = wb.active
    ws.title = "Start"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT
    widths(ws, {"A": 3, "B": 30, "C": 78})

    ws["B1"] = f"Finanzplaner {YEAR}"
    ws["B1"].font = font(24, True, ACCENT)
    ws.row_dimensions[1].height = 40
    ws["B2"] = "Haushaltsbuch · Budget · Sparziele · Schulden-Plan"
    ws["B2"].font = font(11, color=GREY)

    ws["B4"] = "Einstellungen"
    ws["B4"].font = font(12, True, ACCENT)
    ws["B4"].border = BOTTOM
    ws["C4"].border = BOTTOM
    ws["B6"] = "Jahr"
    ws["C6"] = YEAR
    ws["B7"] = "Kontostand am 01.01. (€)"
    ws["C7"] = 1250 if demo else 0
    ws["C7"].number_format = EUR
    for ref in ("C6", "C7"):
        ws[ref].fill = fill(INPUT)
        ws[ref].border = BOX
        ws[ref].font = font(11, True, "0000FF")
        ws[ref].alignment = Alignment(horizontal="left")
    ws["B6"].font = ws["B7"].font = font(11)

    ws["B9"] = "Legende"
    ws["B9"].font = font(12, True, ACCENT)
    ws["B9"].border = BOTTOM
    ws["C9"].border = BOTTOM
    ws["B10"] = "Gelbe Felder"
    ws["B10"].fill = fill(INPUT)
    ws["B10"].border = BOX
    ws["C10"] = "Hier trägst du etwas ein."
    ws["B11"] = "Weiße Felder"
    ws["B11"].border = BOX
    ws["C11"] = "Rechnen automatisch – bitte nicht überschreiben."

    ws["B13"] = "So legst du los (10 Minuten)"
    ws["B13"].font = font(12, True, ACCENT)
    ws["B13"].border = BOTTOM
    ws["C13"].border = BOTTOM
    steps = [
        ("1. Einstellungen", "Trage oben das Jahr und deinen Kontostand am 1. Januar ein."),
        ("2. Kategorien", "Passe die Kategorien an und lege für jede ein Monatsbudget fest."),
        ("3. Buchungen", "Trage jede Einnahme und Ausgabe ein: Datum, Kategorie (Auswahlliste), Betrag. "
                         "Beträge immer positiv – die Kategorie entscheidet, ob Einnahme oder Ausgabe."),
        ("4. Monatsübersicht", "Zeigt automatisch jeden Monat pro Kategorie. Rot = Monatsbudget überschritten."),
        ("5. Dashboard", "Dein Jahr auf einen Blick: Einnahmen, Ausgaben, Sparquote, größte Ausgaben."),
        ("6. Sparziele & Schulden", "Rechnet aus, wann du dein Ziel erreichst bzw. schuldenfrei bist."),
    ]
    for i, (k, v) in enumerate(steps):
        r = 14 + i
        ws.cell(row=r, column=2, value=k).font = font(10, True)
        c = ws.cell(row=r, column=3, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
        ws.row_dimensions[r].height = 28 if len(v) > 80 else 16

    ws["B21"] = "Tipp"
    ws["B21"].font = font(10, True, ACCENT)
    ws["B21"].alignment = Alignment(vertical="top")
    ws["C21"] = ("10 Minuten pro Woche reichen. Trag am besten jeden Sonntag die Buchungen der Woche nach – "
                 "nach einem Monat weißt du genau, wohin dein Geld fließt.")
    ws["C21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[21].height = 28
    ws["B22"] = "Wichtig"
    ws["B22"].font = font(10, True, ACCENT)
    ws["B22"].alignment = Alignment(vertical="top")
    ws["C22"] = ("Wenn du eine Kategorie umbenennst, ändere sie auch in deinen bisherigen Buchungen – "
                 "sonst werden diese nicht mehr mitgezählt (Spalte „Art“ zeigt dann „unbekannte Kategorie“).")
    ws["C22"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[22].height = 28

    ws["B24"] = "Funktioniert in Microsoft Excel, Google Sheets und LibreOffice."
    ws["B24"].font = font(9, color=GREY, italic=True)
    if demo:
        ws["B25"] = "DEMO-VERSION mit Beispieldaten"
        ws["B25"].font = font(10, True, EXPENSE)
    style_all(ws, 25, 3)
    landscape(ws)


# --------------------------------------------------------------------------- Kategorien
def build_categories(wb, demo):
    ws = wb.create_sheet("Kategorien")
    ws.sheet_properties.tabColor = "8FB9B5"
    title(ws, "Kategorien & Monatsbudget",
          "Namen frei änderbar. Leere Zeilen kannst du für eigene Kategorien nutzen.")
    widths(ws, {"A": 34, "B": 12, "C": 20, "D": 4, "E": 60})
    header_row(ws, 4, ["Kategorie", "Art", "Monatsbudget (€)"])

    demo_budget = {
        "Gehalt / Lohn": 2150, "Nebenjob": 200, "Verkäufe (Kleinanzeigen etc.)": 30,
        "Miete & Nebenkosten": 720, "Strom & Internet": 125, "Lebensmittel": 320,
        "Handy & Abos": 45, "Mobilität (Auto/ÖPNV)": 63, "Versicherungen": 17,
        "Freizeit & Ausgehen": 150, "Kleidung": 60, "Gesundheit & Drogerie": 40,
        "Geschenke": 30, "Urlaub": 80, "Sonstiges": 40,
    }
    for r in range(INC_FIRST, EXP_LAST + 1):
        is_inc = r <= INC_LAST
        names = INCOME_CATS if is_inc else EXPENSE_CATS
        idx = r - (INC_FIRST if is_inc else EXP_FIRST)
        name = names[idx] if idx < len(names) else None
        a = ws.cell(row=r, column=1, value=name)
        b = ws.cell(row=r, column=2, value="Einnahme" if is_inc else "Ausgabe")
        c = ws.cell(row=r, column=3, value=demo_budget.get(name) if demo and name else None)
        a.fill = c.fill = fill(INPUT)
        a.font = c.font = font(10, color="0000FF")
        b.fill = fill(INCOME_LIGHT if is_inc else EXPENSE_LIGHT)
        b.font = font(9, True, INCOME if is_inc else EXPENSE)
        b.alignment = Alignment(horizontal="center")
        c.number_format = EUR
        for cell in (a, b, c):
            cell.border = BOX

    ws["E4"] = "Hinweise"
    ws["E4"].font = font(10, True, ACCENT)
    notes = [
        f"Zeilen {INC_FIRST}–{INC_LAST}: Einnahmen (max. {INC_LAST - INC_FIRST + 1} Kategorien)",
        f"Zeilen {EXP_FIRST}–{EXP_LAST}: Ausgaben (max. {EXP_LAST - EXP_FIRST + 1} Kategorien)",
        "Das Monatsbudget ist dein Plan. Die Monatsübersicht färbt Monate rot,",
        "in denen du eine Ausgaben-Kategorie überzogen hast.",
        "Faustregel 50/30/20: 50 % Fixkosten, 30 % Freizeit & Wünsche, 20 % Sparen.",
    ]
    for i, n in enumerate(notes):
        ws.cell(row=5 + i, column=5, value=n).font = font(9, color=GREY)
    ws.freeze_panes = "A5"
    style_all(ws, EXP_LAST, 5)
    landscape(ws)


# --------------------------------------------------------------------------- Buchungen
def build_bookings(wb, demo):
    ws = wb.create_sheet("Buchungen")
    ws.sheet_properties.tabColor = "8FB9B5"
    title(ws, "Buchungen",
          "Jede Einnahme und Ausgabe in eine Zeile. Beträge immer positiv – die Kategorie bestimmt, "
          "ob es eine Einnahme oder Ausgabe ist.")
    widths(ws, {"A": 13, "B": 32, "C": 42, "D": 15, "E": 20})
    header_row(ws, 4, ["Datum", "Kategorie", "Beschreibung", "Betrag (€)", "Art (automatisch)"])

    rows = demo_bookings() if demo else [
        (dt.date(YEAR, 1, 1), "Gehalt / Lohn", "Beispiel – diese Zeile einfach überschreiben", 2100.00)
    ]
    for r in range(B_FIRST, B_LAST + 1):
        i = r - B_FIRST
        if i < len(rows):
            d, cat, desc, amt = rows[i]
            ws.cell(row=r, column=1, value=d)
            ws.cell(row=r, column=2, value=cat)
            ws.cell(row=r, column=3, value=desc)
            ws.cell(row=r, column=4, value=amt)
        ws.cell(row=r, column=1).number_format = DATE
        ws.cell(row=r, column=4).number_format = EUR
        ws.cell(row=r, column=5, value=(
            f'=IF(B{r}="","",IFERROR(INDEX(Kategorien!$B${INC_FIRST}:$B${EXP_LAST},'
            f'MATCH(B{r},Kategorien!$A${INC_FIRST}:$A${EXP_LAST},0)),"unbekannte Kategorie"))'
        ))
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.fill = fill(INPUT)
            c.font = font(10, color="0000FF")
            c.border = BOX
        e = ws.cell(row=r, column=5)
        e.font = font(9, color=GREY)
        e.alignment = Alignment(horizontal="center")
        e.border = BOX

    dv_cat = DataValidation(type="list", formula1=f"=Kategorien!$A${INC_FIRST}:$A${EXP_LAST}",
                            allow_blank=True, showErrorMessage=True, errorStyle="warning",
                            errorTitle="Unbekannte Kategorie",
                            error="Bitte eine Kategorie aus der Liste wählen (Tab „Kategorien“).")
    dv_cat.add(f"B{B_FIRST}:B{B_LAST}")
    dv_amt = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            allow_blank=True, showErrorMessage=True, errorTitle="Betrag positiv eingeben",
                            error="Beträge immer positiv eingeben – die Kategorie entscheidet, "
                                  "ob Einnahme oder Ausgabe.")
    dv_amt.add(f"D{B_FIRST}:D{B_LAST}")
    dv_date = DataValidation(type="date", operator="between",
                             formula1="DATE(Start!$C$6,1,1)", formula2="DATE(Start!$C$6,12,31)",
                             allow_blank=True, showErrorMessage=True, errorStyle="warning",
                             errorTitle="Anderes Jahr?",
                             error="Das Datum liegt nicht im eingestellten Jahr (Tab „Start“) "
                                   "und wird nicht mitgezählt.")
    dv_date.add(f"A{B_FIRST}:A{B_LAST}")
    for dv in (dv_cat, dv_amt, dv_date):
        ws.add_data_validation(dv)

    rng = f"E{B_FIRST}:E{B_LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'E{B_FIRST}="Einnahme"'],
                                                   font=Font(name=FONT, color=INCOME, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'E{B_FIRST}="Ausgabe"'],
                                                   font=Font(name=FONT, color=EXPENSE, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'E{B_FIRST}="unbekannte Kategorie"'],
                                                   font=Font(name=FONT, color="FFFFFF", bold=True),
                                                   fill=fill(EXPENSE)))
    ws.freeze_panes = f"A{B_FIRST}"
    ws.auto_filter.ref = f"A4:E{B_LAST}"
    landscape(ws)


def demo_bookings():
    rnd = random.Random(42)
    out = []

    def add(m, day, cat, desc, amt):
        out.append((dt.date(YEAR, m, day), cat, desc, round(amt, 2)))

    for m in range(1, 10):  # Januar bis September
        add(m, 1, "Gehalt / Lohn", "Gehalt", 2150.00)
        add(m, 3, "Miete & Nebenkosten", "Warmmiete", 720.00)
        add(m, 5, "Strom & Internet", "Strom Abschlag", 85.00)
        add(m, 5, "Strom & Internet", "Internet", 39.99)
        add(m, 7, "Handy & Abos", "Handyvertrag", 15.99)
        add(m, 7, "Handy & Abos", "Streaming", 13.99)
        add(m, 7, "Handy & Abos", "Musik-Abo", 10.99)
        add(m, 2, "Mobilität (Auto/ÖPNV)", "Deutschlandticket", 63.00)
        add(m, 10, "Versicherungen", "Haftpflicht & Hausrat", 16.30)
        if rnd.random() < 0.7:
            add(m, 15, "Nebenjob", "Minijob", rnd.choice([180, 220, 260, 310]))
        if rnd.random() < 0.4:
            add(m, rnd.randint(8, 25), "Verkäufe (Kleinanzeigen etc.)", "Kleinanzeigen-Verkauf",
                rnd.choice([15, 25, 40, 65, 90]))
        for _ in range(rnd.randint(6, 9)):
            add(m, rnd.randint(1, 28), "Lebensmittel",
                rnd.choice(["Wocheneinkauf", "Supermarkt", "Discounter", "Bäcker", "Markt"]),
                rnd.uniform(14, 72))
        for _ in range(rnd.randint(3, 6)):
            add(m, rnd.randint(1, 28), "Freizeit & Ausgehen",
                rnd.choice(["Kino", "Essen gehen", "Bar mit Freunden", "Fitnessstudio", "Konzert", "Café"]),
                rnd.uniform(8, 58))
        for _ in range(rnd.randint(0, 2)):
            add(m, rnd.randint(1, 28), "Kleidung", rnd.choice(["Schuhe", "Jacke", "T-Shirts", "Hose"]),
                rnd.uniform(19, 89))
        for _ in range(rnd.randint(2, 3)):
            add(m, rnd.randint(1, 28), "Gesundheit & Drogerie", rnd.choice(["Drogerie", "Apotheke"]),
                rnd.uniform(4, 28))
        if rnd.random() < 0.35:
            add(m, rnd.randint(1, 28), "Geschenke", "Geburtstagsgeschenk", rnd.uniform(15, 45))
        for _ in range(rnd.randint(1, 2)):
            add(m, rnd.randint(1, 28), "Sonstiges", rnd.choice(["Haushalt", "Friseur", "Porto", "Bastelkram"]),
                rnd.uniform(5, 38))
    add(7, 14, "Urlaub", "Ferienwohnung Ostsee", 640.00)
    add(8, 2, "Urlaub", "Bahntickets", 89.90)
    out.sort(key=lambda t: t[0])
    return out


# --------------------------------------------------------------------------- Monatsuebersicht
def build_monthly(wb):
    ws = wb.create_sheet("Monatsübersicht")
    ws.sheet_properties.tabColor = ACCENT
    ws["A1"] = '="Monatsübersicht "&Start!C6'
    ws["A1"].font = font(18, True, ACCENT)
    ws.row_dimensions[1].height = 30
    ws["A2"] = "Rechnet automatisch aus deinen Buchungen. Rot hinterlegt = Monatsbudget überschritten."
    ws["A2"].font = font(10, color=GREY, italic=True)
    widths(ws, {"A": 30, "B": 13, **{get_column_letter(c): 10.5 for c in range(3, 15)},
                "O": 13, "P": 12, "Q": 13, "R": 14, "S": 4})

    # Zeile 3: Monatsnummern als Hilfswerte fuer die Formeln (ausgeblendet)
    for i in range(12):
        ws.cell(row=3, column=3 + i, value=i + 1)
    ws.row_dimensions[3].hidden = True
    header_row(ws, 4, ["Kategorie", "Budget / Monat", *MONTHS, "Summe", "Ø / Monat",
                       "Budget bisher", "Abweichung\n(+ = gut)"])

    def section(row, label, color, light):
        for col in range(1, 19):
            ws.cell(row=row, column=col).fill = fill(light)
        c = ws.cell(row=row, column=1, value=label)
        c.font = font(10, True, color)

    def cat_rows(first, last, src_first, is_income):
        for r in range(first, last + 1):
            src = src_first + (r - first)
            ws.cell(row=r, column=1, value=f'=IF(Kategorien!A{src}="","",Kategorien!A{src})')
            ws.cell(row=r, column=2, value=f'=IF($A{r}="",0,N(Kategorien!C{src}))')
            for m in range(12):
                col = get_column_letter(3 + m)
                ws.cell(row=r, column=3 + m, value=(
                    f'=IF($A{r}="",0,SUMIFS({BK_AMT},{BK_CAT},$A{r},'
                    f'{BK_DATE},">="&DATE(Start!$C$6,{col}$3,1),'
                    f'{BK_DATE},"<"&DATE(Start!$C$6,{col}$3+1,1)))'
                ))
            ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})")
            ws.cell(row=r, column=16, value=f'=IF($A{r}="","",O{r}/MAX(1,$B${MO_ACTIVE}))')
            ws.cell(row=r, column=17, value=f'=IF(OR($A{r}="",$B{r}=0),"",B{r}*$B${MO_ACTIVE})')
            diff = f"O{r}-Q{r}" if is_income else f"Q{r}-O{r}"
            ws.cell(row=r, column=18, value=f'=IF(OR($A{r}="",$B{r}=0),"",{diff})')
            if not is_income:
                # Sortierhilfe fuer die Top-Ausgaben im Dashboard (kleiner Zuschlag loest Gleichstaende)
                ws.cell(row=r, column=19, value=f'=IF($A{r}="",-1,O{r}+(100-ROW())/1000000)')
            for col in range(1, 19):
                c = ws.cell(row=r, column=col)
                c.border = BOX
                c.font = font(10, bold=(col == 15))
                if col >= 2:
                    c.number_format = EUR

    def sum_row(row, first, last, label, color):
        ws.cell(row=row, column=1, value=label)
        for col in range(2, 16):
            L = get_column_letter(col)
            ws.cell(row=row, column=col, value=f"=SUM({L}{first}:{L}{last})")
        ws.cell(row=row, column=16, value=f"=O{row}/MAX(1,$B${MO_ACTIVE})")
        ws.cell(row=row, column=17, value=f"=SUM(Q{first}:Q{last})")
        for col in range(1, 19):
            c = ws.cell(row=row, column=col)
            c.font = font(10, True, color)
            c.number_format = EUR
            c.border = Border(top=Side(style="medium", color=color), bottom=thin)

    section(5, "EINNAHMEN", INCOME, INCOME_LIGHT)
    cat_rows(MO_INC_FIRST, MO_INC_LAST, INC_FIRST, True)
    sum_row(MO_INC_SUM, MO_INC_FIRST, MO_INC_LAST, "Summe Einnahmen", INCOME)
    section(16, "AUSGABEN", EXPENSE, EXPENSE_LIGHT)
    cat_rows(MO_EXP_FIRST, MO_EXP_LAST, EXP_FIRST, False)
    sum_row(MO_EXP_SUM, MO_EXP_FIRST, MO_EXP_LAST, "Summe Ausgaben", EXPENSE)

    ws.cell(row=MO_NET, column=1, value="Überschuss (Einnahmen − Ausgaben)")
    ws.cell(row=MO_RATE, column=1, value="Sparquote")
    ws.cell(row=MO_BAL, column=1, value="Kontostand am Monatsende")
    for col in range(3, 16):
        L = get_column_letter(col)
        ws.cell(row=MO_NET, column=col, value=f"={L}{MO_INC_SUM}-{L}{MO_EXP_SUM}").number_format = EUR
        ws.cell(row=MO_RATE, column=col,
                value=f"=IF({L}{MO_INC_SUM}=0,0,{L}{MO_NET}/{L}{MO_INC_SUM})").number_format = PCT
        if col <= 14:
            ws.cell(row=MO_BAL, column=col,
                    value=f"=Start!$C$7+SUM($C${MO_NET}:{L}{MO_NET})").number_format = EUR
    for r in (MO_NET, MO_RATE, MO_BAL):
        for col in range(1, 16):
            c = ws.cell(row=r, column=col)
            c.font = font(10, True, ACCENT)
            c.fill = fill(ACCENT_LIGHT)
            c.border = BOX

    ws.cell(row=MO_ACTIVE, column=1, value="Monate mit Buchungen").font = font(9, color=GREY)
    c = ws.cell(row=MO_ACTIVE, column=2,
                value=f"=SUMPRODUCT(--((C{MO_INC_SUM}:N{MO_INC_SUM}+C{MO_EXP_SUM}:N{MO_EXP_SUM})>0))")
    c.font = font(9, True, GREY)
    c.alignment = Alignment(horizontal="left")
    ws.cell(row=MO_ACTIVE + 1, column=1,
            value="„Ø / Monat“ und „Budget bisher“ beziehen sich auf die Monate mit Buchungen.").font = \
        font(9, color=GREY, italic=True)

    ws.conditional_formatting.add(
        f"C{MO_EXP_FIRST}:N{MO_EXP_LAST}",
        FormulaRule(formula=[f"AND($B{MO_EXP_FIRST}>0,C{MO_EXP_FIRST}>$B{MO_EXP_FIRST})"],
                    fill=fill("F8C9C4"), font=Font(name=FONT, color=EXPENSE, bold=True)))
    for rng in (f"R{MO_INC_FIRST}:R{MO_INC_LAST}", f"R{MO_EXP_FIRST}:R{MO_EXP_LAST}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                      font=Font(name=FONT, color=INCOME, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                      font=Font(name=FONT, color=EXPENSE, bold=True)))
    ws.conditional_formatting.add(f"C{MO_NET}:O{MO_NET}",
                                  CellIsRule(operator="lessThan", formula=["0"],
                                             font=Font(name=FONT, color=EXPENSE, bold=True)))
    for first, last in ((MO_INC_FIRST, MO_INC_LAST), (MO_EXP_FIRST, MO_EXP_LAST)):
        ws.conditional_formatting.add(f"B{first}:R{last}", FormulaRule(
            formula=[f'$A{first}=""'], font=Font(name=FONT, color="FFFFFF")))
    ws.column_dimensions["S"].hidden = True
    ws.freeze_panes = "C5"
    ws.protection.sheet = True  # ohne Passwort – schuetzt nur vor versehentlichem Ueberschreiben
    style_all(ws, MO_ACTIVE + 1, 19)
    landscape(ws)


# --------------------------------------------------------------------------- Dashboard
def build_dashboard(wb):
    ws = wb.create_sheet("Dashboard", 1)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT
    ws["A1"] = '="Dein Finanzjahr "&Start!C6'
    ws["A1"].font = font(20, True, ACCENT)
    ws.row_dimensions[1].height = 34
    ws["A2"] = "Alles hier rechnet automatisch aus deinen Buchungen."
    ws["A2"].font = font(10, color=GREY, italic=True)
    widths(ws, {"A": 2, **{get_column_letter(c): 13 for c in range(2, 22)}})

    mo = "Monatsübersicht"
    inc, exp, net = f"{mo}!O{MO_INC_SUM}", f"{mo}!O{MO_EXP_SUM}", f"{mo}!O{MO_NET}"
    active = f"{mo}!B{MO_ACTIVE}"

    def tiles(row, items):
        for i, (label, formula, fmt, color) in enumerate(items):
            col = 2 + i * 2
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
            ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
            lab = ws.cell(row=row, column=col, value=label)
            lab.font = font(9, True, GREY)
            val = ws.cell(row=row + 1, column=col, value=formula)
            val.font = font(16, True, color)
            val.number_format = fmt
            for r in (row, row + 1):
                for cc in (col, col + 1):
                    cell = ws.cell(row=r, column=cc)
                    cell.fill = fill(ACCENT_LIGHT)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row].height = 20
            ws.row_dimensions[row + 1].height = 30

    tiles(4, [
        ("EINNAHMEN", f"={inc}", EUR0, INCOME),
        ("AUSGABEN", f"={exp}", EUR0, EXPENSE),
        ("ÜBERSCHUSS", f"={net}", EUR0, ACCENT),
        ("SPARQUOTE", f"=IF({inc}=0,0,{net}/{inc})", PCT, ACCENT),
        ("KONTOSTAND AKTUELL", f"=Start!C7+{net}", EUR0, ACCENT),
    ])
    tiles(7, [
        ("Ø EINNAHMEN / MONAT", f"={inc}/MAX(1,{active})", EUR0, INCOME),
        ("Ø AUSGABEN / MONAT", f"={exp}/MAX(1,{active})", EUR0, EXPENSE),
        ("Ø ÜBERSCHUSS / MONAT", f"={net}/MAX(1,{active})", EUR0, ACCENT),
        ("MONATE ERFASST", f"={active}", "0", ACCENT),
        ("KATEGORIEN ÜBER BUDGET", f'=COUNTIF({mo}!R{MO_EXP_FIRST}:R{MO_EXP_LAST},"<0")', "0", EXPENSE),
    ])

    # Monatsverlauf (Datenbasis fuer Diagramm 1)
    ws["B11"] = "Monatsverlauf"
    ws["B11"].font = font(12, True, ACCENT)
    header_row(ws, 12, ["Monat", "Einnahmen", "Ausgaben", "Überschuss"], start_col=2)
    for i, name in enumerate(MONTHS):
        r = 13 + i
        L = get_column_letter(3 + i)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=f"={mo}!{L}{MO_INC_SUM}")
        ws.cell(row=r, column=4, value=f"={mo}!{L}{MO_EXP_SUM}")
        ws.cell(row=r, column=5, value=f"={mo}!{L}{MO_NET}")
        for col in range(2, 6):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = font(10)
            if col > 2:
                c.number_format = EUR0
    ws.conditional_formatting.add("E13:E24", CellIsRule(operator="lessThan", formula=["0"],
                                                        font=Font(name=FONT, color=EXPENSE, bold=True)))

    ch = BarChart()
    ch.type = "col"
    ch.title = "Einnahmen vs. Ausgaben"
    ch.y_axis.title = None
    ch.y_axis.numFmt = '#,##0 "€"'
    ch.y_axis.majorGridlines = None
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=12, max_row=24), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=13, max_row=24))
    ch.series[0].graphicalProperties.solidFill = INCOME
    ch.series[0].graphicalProperties.line.solidFill = INCOME
    ch.series[1].graphicalProperties.solidFill = EXPENSE
    ch.series[1].graphicalProperties.line.solidFill = EXPENSE
    ch.gapWidth = 60
    ch.legend.position = "b"
    ch.height, ch.width = 7.5, 17
    ws.add_chart(ch, "G11")

    # Top-8-Ausgaben (Datenbasis fuer Diagramm 2)
    ws["B27"] = "Deine größten Ausgaben"
    ws["B27"].font = font(12, True, ACCENT)
    header_row(ws, 28, ["Kategorie", "", "Summe", "Anteil"], start_col=2)
    ws.merge_cells("B28:C28")
    sort = f"{mo}!$S${MO_EXP_FIRST}:$S${MO_EXP_LAST}"
    names = f"{mo}!$A${MO_EXP_FIRST}:$A${MO_EXP_LAST}"
    sums = f"{mo}!$O${MO_EXP_FIRST}:$O${MO_EXP_LAST}"
    for k in range(1, 9):
        r = 28 + k
        pos = f"MATCH(LARGE({sort},{k}),{sort},0)"
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(row=r, column=2, value=f'=IF(INDEX({sums},{pos})<=0,"",INDEX({names},{pos}))')
        ws.cell(row=r, column=4, value=f"=MAX(0,INDEX({sums},{pos}))").number_format = EUR0
        ws.cell(row=r, column=5, value=f"=IF({exp}=0,0,D{r}/{exp})").number_format = PCT
        for col in range(2, 6):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = font(10)

    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "Top-Ausgaben nach Kategorie"
    ch2.add_data(Reference(ws, min_col=4, min_row=29, max_row=36), titles_from_data=False)
    ch2.set_categories(Reference(ws, min_col=2, min_row=29, max_row=36))
    ch2.series[0].graphicalProperties.solidFill = ACCENT
    ch2.series[0].graphicalProperties.line.solidFill = ACCENT
    ch2.x_axis.scaling.orientation = "maxMin"  # groesste Ausgabe oben
    ch2.y_axis.numFmt = '#,##0 "€"'
    ch2.y_axis.majorGridlines = None
    ch2.legend = None
    ch2.gapWidth = 50
    ch2.height, ch2.width = 7.5, 17
    ws.add_chart(ch2, "G27")

    ws.protection.sheet = True
    style_all(ws, 36, 11)
    landscape(ws)
    ws.page_setup.fitToHeight = 1


# --------------------------------------------------------------------------- Sparziele
def build_goals(wb, demo):
    ws = wb.create_sheet("Sparziele")
    ws.sheet_properties.tabColor = INCOME
    title(ws, "Sparziele", "Trag dein Ziel, den Betrag und deine monatliche Sparrate ein – "
                           "der Rest rechnet sich von selbst.")
    widths(ws, {"A": 30, "B": 15, "C": 15, "D": 16, "E": 15, "F": 14, "G": 15, "H": 17})
    header_row(ws, 4, ["Ziel", "Zielbetrag (€)", "Schon gespart (€)", "Sparrate / Monat (€)",
                       "Fehlt noch (€)", "Fortschritt", "Monate bis Ziel", "Ziel erreicht ca."])
    data = [
        ("Notgroschen (3 Monatsausgaben)", 4500, 1800, 150),
        ("Urlaub nächstes Jahr", 1500, 300, 100),
        ("Neuer Laptop", 1200, 450, 75),
    ] if demo else [("Notgroschen (Beispiel – überschreiben)", 3000, 0, 100)]
    for r in range(5, 15):
        i = r - 5
        if i < len(data):
            for col, v in enumerate(data[i], start=1):
                ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=5, value=f'=IF(A{r}="","",MAX(0,N(B{r})-N(C{r})))')
        ws.cell(row=r, column=6, value=f'=IF(OR(A{r}="",N(B{r})=0),"",MIN(1,N(C{r})/B{r}))')
        ws.cell(row=r, column=7, value=(
            f'=IF(A{r}="","",IF(E{r}=0,0,IF(N(D{r})<=0,"Rate fehlt",ROUNDUP(E{r}/D{r},0))))'))
        ws.cell(row=r, column=8, value=(
            f'=IF(ISNUMBER(G{r}),IF(G{r}=0,"erreicht!",EDATE(TODAY(),G{r})),"")'))
        for col in range(1, 9):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col <= 4:
                c.fill = fill(INPUT)
                c.font = font(10, color="0000FF")
            else:
                c.font = font(10, bold=(col == 8))
            if col in (2, 3, 4, 5):
                c.number_format = EUR
        ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=r, column=7).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=8).number_format = "MMM YYYY"
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("F5:F14", DataBarRule(start_type="num", start_value=0,
                                                        end_type="num", end_value=1, color="5FA35F"))
    ws["A16"] = "Summe"
    for col, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        c = ws.cell(row=16, column=col, value=f"=SUM({L}5:{L}14)")
        c.number_format = EUR
    for col in range(1, 9):
        ws.cell(row=16, column=col).font = font(10, True, ACCENT)
        ws.cell(row=16, column=col).border = Border(top=Side(style="medium", color=ACCENT))
    ws["A18"] = ("Tipp: Bau zuerst einen Notgroschen von etwa 3 Monatsausgaben auf. "
                 "Richte die Sparrate als Dauerauftrag direkt nach dem Gehaltseingang ein.")
    ws["A18"].font = font(9, color=GREY, italic=True)
    style_all(ws, 18, 8)
    landscape(ws)


# --------------------------------------------------------------------------- Schulden
def build_debts(wb, demo):
    ws = wb.create_sheet("Schulden")
    ws.sheet_properties.tabColor = EXPENSE
    title(ws, "Schulden-Plan", "Trag Restbetrag, Zinssatz und Monatsrate ein – du siehst sofort, "
                               "wann du schuldenfrei bist und wie viel Zinsen du zahlst.")
    widths(ws, {"A": 30, "B": 15, "C": 13, "D": 15, "E": 17, "F": 17, "G": 18})
    header_row(ws, 4, ["Schuld", "Restbetrag (€)", "Zinssatz p.a.", "Monatsrate (€)",
                       "Monate bis schuldenfrei", "Schuldenfrei ca.", "Zinsen gesamt ca. (€)"])
    data = [
        ("Dispokredit", 850, 0.129, 100),
        ("Ratenkredit Auto", 4200, 0.069, 160),
    ] if demo else [("Dispo (Beispiel – überschreiben)", 500, 0.129, 50)]
    for r in range(5, 15):
        i = r - 5
        if i < len(data):
            for col, v in enumerate(data[i], start=1):
                ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=5, value=(
            f'=IF(A{r}="","",IF(N(B{r})<=0,0,IF(N(D{r})<=0,"Rate fehlt",IF(N(C{r})=0,ROUNDUP(B{r}/D{r},0),'
            f'IF(D{r}<=B{r}*C{r}/12,"Rate zu niedrig",ROUNDUP(NPER(C{r}/12,-D{r},B{r}),0))))))'))
        ws.cell(row=r, column=6, value=(
            f'=IF(ISNUMBER(E{r}),IF(E{r}=0,"schuldenfrei!",EDATE(TODAY(),E{r})),"")'))
        ws.cell(row=r, column=7, value=(
            f'=IF(ISNUMBER(E{r}),IF(OR(E{r}=0,N(C{r})=0),0,D{r}*NPER(C{r}/12,-D{r},B{r})-B{r}),"")'))
        for col in range(1, 8):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col <= 4:
                c.fill = fill(INPUT)
                c.font = font(10, color="0000FF")
            else:
                c.font = font(10, bold=(col == 6))
            if col in (2, 4, 7):
                c.number_format = EUR
        ws.cell(row=r, column=3).number_format = "0.00%"
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=6).number_format = "MMM YYYY"
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("E5:E14", FormulaRule(formula=['ISTEXT(E5)'],
                                                        font=Font(name=FONT, color=EXPENSE, bold=True)))
    ws["A16"] = "Summe"
    ws["B16"] = "=SUM(B5:B14)"
    ws["D16"] = "=SUM(D5:D14)"
    ws["G16"] = "=SUM(G5:G14)"
    ws["E16"] = '=IF(COUNT(E5:E14)=0,"",MAX(E5:E14))'
    for col in range(1, 8):
        c = ws.cell(row=16, column=col)
        c.font = font(10, True, ACCENT)
        c.border = Border(top=Side(style="medium", color=ACCENT))
        if col in (2, 4, 7):
            c.number_format = EUR
    ws["E16"].alignment = Alignment(horizontal="center")
    ws["E17"] = "(längste Laufzeit)"
    ws["E17"].font = font(8, color=GREY, italic=True)
    ws["E17"].alignment = Alignment(horizontal="center")
    ws["A19"] = ("Lawinen-Methode: Alle Mindestraten zahlen und jeden zusätzlichen Euro in die Schuld mit dem "
                 "höchsten Zinssatz stecken – das spart am meisten Zinsen.")
    ws["A20"] = ("Die Berechnung geht von gleichbleibender Rate und gleichbleibendem Zins aus und ist eine "
                 "Schätzung – maßgeblich ist dein Kreditvertrag.")
    ws["A19"].font = ws["A20"].font = font(9, color=GREY, italic=True)
    style_all(ws, 20, 7)
    landscape(ws)


def build(demo):
    wb = Workbook()
    build_start(wb, demo)
    build_categories(wb, demo)
    build_bookings(wb, demo)
    build_monthly(wb)
    build_dashboard(wb)  # wird als 2. Tab einsortiert
    build_goals(wb, demo)
    build_debts(wb, demo)
    order = ["Start", "Dashboard", "Buchungen", "Kategorien", "Monatsübersicht", "Sparziele", "Schulden"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.properties.title = f"Finanzplaner {YEAR}"
    wb.properties.creator = "Finanzplaner"
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / (f"Finanzplaner-{YEAR}-DEMO.xlsx" if demo else f"Finanzplaner-{YEAR}.xlsx")
    wb.save(path)
    print(path)


if __name__ == "__main__":
    build(demo=False)
    build(demo=True)
